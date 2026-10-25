"""Create an entirely invented workbook and run v0.1 without API access."""
import argparse
import subprocess
import sys
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import PatternFill
import fyp_feedback_generator as generator

def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=True)
    source = args.output / "synthetic.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Summary"
    headers = generator.CATEGORIES.copy()
    headers[5] = "Summarisation and Re-evaluation"
    sheet.append(["Fictional team role"] + headers)
    for row, color in enumerate(["0092D050", "00FF0000", "FF5FA052"], start=2):
        sheet.cell(row, 1, f"Invented dialogue for software demonstration, row {row}.")
        for column, category in enumerate(generator.CATEGORIES, start=2):
            cell = sheet.cell(row, column,
                f"Synthetic {category} observation {row}; not a participant assessment.")
            cell.fill = PatternFill("solid", fgColor=color)
    workbook.save(source)
    workbook.close()
    subprocess.run([sys.executable, str(Path(generator.__file__)), str(source),
                    "--output", str(args.output / "reports")], check=True)

if __name__ == "__main__":
    main()

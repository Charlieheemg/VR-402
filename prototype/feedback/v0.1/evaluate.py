"""Re-run the frozen v0.1 locally; publish aggregates only, keep drafts private."""
from __future__ import annotations
import argparse
from collections import Counter
import hashlib
import json
from pathlib import Path
import platform
import re
import openpyxl
import fyp_feedback_generator as generator

# Independent header lookup: do not reuse the extractor's normalize function.
HEADERS = {
    "roles & responsibilities", "closed loop communication", "clear messages",
    "knowing one's limitations", "constructive intervention", "knowledge sharing",
    "summarisation & re-evaluation", "summarization and re-evaluation",
    "summarisation and re-evaluation",
}

def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()

def evaluate(inputs, output):
    paths = sorted(inputs.glob("*.xlsx"))
    if not paths:
        raise ValueError("No .xlsx workbooks found")
    output.mkdir(parents=True, exist_ok=True)
    groups = {}
    private_manifest = []
    records_by_path = {}
    for path in paths:
        records = generator.extract(path)
        records_by_path[path] = records
        workbook = openpyxl.load_workbook(path, data_only=True)
        independent_count = 0
        text_mismatches = 0
        unknown_fills = Counter()
        for sheet in workbook:
            columns = [c.column for c in sheet[1]
                       if str(c.value).strip().lower() in HEADERS]
            independent_count += sum(
                c.value is not None and bool(str(c.value).strip())
                for column in columns
                for row in range(2, sheet.max_row + 1)
                for c in [sheet.cell(row, column)])
        for record in records:
            cell = workbook[record["sheet"]][record["cell"]]
            text_mismatches += str(cell.value) != record["assessment_text"]
            if record["polarity"] == "unverified":
                fill = cell.fill.fgColor
                unknown_fills.update([str(fill.rgb) if fill.type == "rgb" else str(fill.type)])
        workbook.close()
        # Filename groups are provenance labels, not inferred performance cohorts.
        group = "AI" if path.name == "excel_final.xlsx" else (
            "CT" if path.name.startswith("KCT_") else
            "NN" if path.name.endswith("_NN.xlsx") else "other")
        stats = groups.setdefault(group, {"workbooks": 0, "annotation_cells": 0,
            "count_mismatches": 0, "exact_cell_text_mismatches": 0,
            "polarity": Counter(), "unverified_fill_codes": Counter()})
        stats["workbooks"] += 1
        stats["annotation_cells"] += len(records)
        stats["count_mismatches"] += len(records) != independent_count
        stats["exact_cell_text_mismatches"] += text_mismatches
        stats["polarity"].update(r["polarity"] for r in records)
        stats["unverified_fill_codes"].update(unknown_fills)
        private_manifest.append({"file": path.name, "sha256": digest(path),
                                 "records": len(records)})

    # Repeat the packaged design: AI workbook and one consultant pair, not 21 reports.
    ai = [p for p in paths if p.name == "excel_final.xlsx"]
    ct = [p for p in paths if p.name.startswith("KCT_")]
    nn = [p for p in paths if p.name.endswith("_NN.xlsx")]
    # Choose the first consultant pair by their filename token without publishing it.
    pairs = [(left, right) for left in ct for right in nn
             if (left_token := re.search(r"\d+", left.stem))
             and (right_token := re.search(r"\d+", right.stem))
             and left_token.group() == right_token.group()]
    chosen = ai[:1] + (list(pairs[0]) if pairs else [])
    report_checks = []
    for source in chosen:
        name = source.name
        records = records_by_path[source]
        drafts = output / source.stem
        drafts.mkdir(exist_ok=True)
        generator.build_reports(records, drafts, 2)
        repeat = drafts / "repeat"
        repeat.mkdir(exist_ok=True)
        generator.build_reports(records, repeat, 2)
        selected = [r for category in generator.CATEGORIES
                    for r in [x for x in records if x["category"] == category][:2]]
        learner = (drafts / "learner_draft.md").read_text(encoding="utf-8")
        instructor = (drafts / "instructor_draft.md").read_text(encoding="utf-8")
        pointers = [f"{r['workbook']} / {r['sheet']}!{r['cell']}" for r in selected]
        report_checks.append({"source_group": "AI" if name == "excel_final.xlsx" else
                             "CT" if name.startswith("KCT_") else "NN",
            "extracted": len(records), "displayed": len(selected),
            "not_displayed": len(records) - len(selected),
            "all_categories_visible": all("## " + c in learner and
                "### " + c in instructor for c in generator.CATEGORIES),
            "displayed_pointers_in_both": all(p in learner and p in instructor for p in pointers),
            "displayed_annotations_exact": all(r["assessment_text"] in learner and
                r["assessment_text"] in instructor for r in selected),
            "repeat_identical": all((drafts / f).read_bytes() == (repeat / f).read_bytes()
                for f in ["learner_draft.md", "instructor_draft.md"]),
            "source_record_serialization_repeat_identical":
                json.dumps(records, ensure_ascii=False) == json.dumps(generator.extract(source), ensure_ascii=False)})

    result = {"version": "0.1", "mode": "deterministic/local",
        "python": platform.python_version(), "openpyxl": openpyxl.__version__,
        "generator_sha256": digest(Path(generator.__file__)),
        "workbooks": len(paths), "groups": groups, "report_checks": report_checks,
        "count_check_scope": "All populated cells under independently recognised headers; identical header-layout assumption remains.",
        "fidelity_scope": "Exact extracted cell text and displayed annotation/pointer presence only; no event meaning or clinical review.",
        "external_calls": False, "human_or_learning_evaluation": False}
    (output / "aggregate.json").write_text(json.dumps(result, indent=2) + "\n", encoding="utf-8")
    (output / "private_manifest.json").write_text(json.dumps(private_manifest, indent=2), encoding="utf-8")
    return result

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--inputs", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True,
                        help="Use an ignored private/ directory; drafts contain source text")
    args = parser.parse_args()
    print(json.dumps(evaluate(args.inputs, args.output), indent=2))

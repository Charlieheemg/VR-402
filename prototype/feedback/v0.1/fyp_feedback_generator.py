#!/usr/bin/env python3
"""FYP Part B: evidence-preserving learner/instructor feedback prototype.

This script DOES NOT score clinical performance, establish gold standards, or
validate educational effectiveness. It reads existing rubric-labelled worksheets,
creates auditable evidence records, and generates draft reports for expert review.

Optional LLM drafting requires explicit --use-llm AND --allow-external-data.
External API transfer of clinical transcripts requires appropriate authorisation.
"""
from __future__ import annotations
import argparse, csv, json, os, re, sys, urllib.request, urllib.error
from collections import Counter, defaultdict
from pathlib import Path
from openpyxl import load_workbook

CATEGORIES = ["Roles & Responsibilities", "Closed Loop Communication", "Clear Messages", "Knowing One's Limitations", "Constructive Intervention", "Summarisation & Re-evaluation", "Knowledge Sharing"]
ALIASES = {"roles & responsibilities":CATEGORIES[0],"closed loop communication":CATEGORIES[1],"clear messages":CATEGORIES[2],"knowing one's limitations":CATEGORIES[3],"constructive intervention":CATEGORIES[4],"summarization and re-evaluation":CATEGORIES[5],"summarisation and re-evaluation":CATEGORIES[5],"knowledge sharing":CATEGORIES[6]}
TARGETS = {
 CATEGORIES[0]:"Make role allocation, delegation and hand-offs explicit.",
 CATEGORIES[1]:"Use directed requests, explicit acknowledgements and completion confirmations where appropriate.",
 CATEGORIES[2]:"Communicate an identifiable receiver, action and relevant details clearly.",
 CATEGORIES[3]:"Disclose uncertainty and seek assistance or escalation when needed.",
 CATEGORIES[4]:"Speak up about observed problems and suggest corrections constructively.",
 CATEGORIES[5]:"Summarise patient/team status and re-evaluate the plan when appropriate.",
 CATEGORIES[6]:"Share relevant and timely information with the team."
}

def normalize(s):
 return ALIASES.get(str(s).strip().lower())

def polarity(cell):
 # Formatting may encode polarity; colour interpretations are provisional.
 c=cell.fill.fgColor
 if cell.fill.patternType!='solid' or c.type!='rgb': return 'unverified'
 v=(c.rgb or '')[-6:].upper()
 if v in ('00FF00','92D050','C6EFCE','70AD47','00B050'):return 'positive_color'
 if v in ('FF0000','FFC7CE','F4CCCC','FF6666','C00000'):return 'negative_color'
 return 'unverified'

def extract(path):
 wb=load_workbook(path,data_only=True)
 rows=[]
 for ws in wb:
  headers={c.column:str(c.value).strip() for c in ws[1] if c.value is not None}
  cats={col:normalize(v) for col,v in headers.items() if normalize(v)}
  speakercols={col:v for col,v in headers.items() if col not in cats and v.lower() not in ('n/a','na')}
  for row in range(2,ws.max_row+1):
   dialogue=[{'role':label,'text':str(ws.cell(row,col).value),'cell':ws.cell(row,col).coordinate} for col,label in speakercols.items() if ws.cell(row,col).value is not None and str(ws.cell(row,col).value).strip()]
   for col,cat in cats.items():
    cell=ws.cell(row,col)
    if cell.value is None or not str(cell.value).strip():continue
    # Include limited surrounding dialogue, never infer a speaker from row alone.
    context=[]
    for r in range(max(2,row-2),min(ws.max_row,row+2)+1):
     for scol,label in speakercols.items():
      v=ws.cell(r,scol).value
      if v is not None and str(v).strip():context.append({'role':label,'text':str(v),'cell':ws.cell(r,scol).coordinate})
    rows.append({'id':f'{path.stem}:{ws.title}!{cell.coordinate}','workbook':path.name,'sheet':ws.title,'cell':cell.coordinate,'row':row,'category':cat,'assessment_text':str(cell.value),'polarity':polarity(cell),'same_row_dialogue':dialogue,'nearby_dialogue':context,'attribution':'team_or_unverified','review_status':'not_verified'})
 return rows

def md(s):return str(s).replace('|','\\|').replace('\n',' ')

def build_reports(records,out,limit):
 bycat=defaultdict(list)
 for r in records:bycat[r['category']].append(r)
 # Order is fixed by rubric, NOT a clinically validated ranking.
 learner=['# Learner-facing feedback — DRAFT','', '> Concept prototype. All observations are team/role-context until attribution is reviewed. Not clinically validated.','', '## Learning overview','', 'Seven competencies are retained. No clinical priority score or individual competence rating is inferred.','']
 instructor=['# Instructor-facing feedback — DRAFT','', '> Concept prototype for group debriefing. Counts below are annotation records, not independent behaviour opportunities or competence scores.','', '## Competency coverage','', '| Competency | Annotation records |','|---|---:|']
 for cat in CATEGORIES:instructor.append(f'| {cat} | {len(bycat[cat])} |')
 instructor+=['','## Discussion evidence by competency','']
 for cat in CATEGORIES:
  learner += [f'## {cat}','',f'**Expected behaviour (provisional):** {TARGETS[cat]}','']
  instructor += [f'### {cat}','',f'**Discussion target (provisional):** {TARGETS[cat]}','']
  items=bycat[cat][:limit]
  if not items:
   learner+=['No annotation recorded in this category. This does **not** imply satisfactory performance.','']
   instructor+=['No annotation recorded; not evidence of absence.','']
  for r in items:
   src=f"{r['workbook']} / {r['sheet']}!{r['cell']}"
   ev='; '.join(f"{x['role']} ({x['cell']}): {x['text']}" for x in r['nearby_dialogue'][:6]) or 'No nearby transcript found'
   learner += [f'**Observed annotation:** {r["assessment_text"]}',f'- Source: `{src}`',f'- Provisional polarity: `{r["polarity"]}` (verify colour conventions)',f'- Nearby transcript (not confirmed complete event span): {ev}', '- **Reflection:** What did the team do here, and how does it compare with the expected behaviour?', '- **Possible next practice:** Rehearse the target behaviour in a comparable scenario; instructor to specify and approve the action.', '- **Review flag:** Original event span, attribution, rubric and recommendation require expert/source checks.','']
   instructor += [f'- **{src}** — {r["assessment_text"]} (polarity `{r["polarity"]}`; verify).',f'  - Context: {ev}', '  - Suggested discussion: What happened, how was it understood, and what could the team do differently?', '  - No individual blame or clinical consequence inferred.','']
 learner+=['## Learner-owned practice goal','', 'After discussion with the educator: identify one approved practice goal and how to check it next time. This section is intentionally not automatically ranked.','']
 instructor+=['## Debrief plan','', '1. Invite team self-assessment. 2. Examine one source-verified strength and one source-verified gap. 3. Agree on a practice goal. 4. Record unresolved interpretations.','']
 (out/'learner_draft.md').write_text('\n'.join(learner),encoding='utf-8')
 (out/'instructor_draft.md').write_text('\n'.join(instructor),encoding='utf-8')

def api_draft(records,model,timeout):
 key=os.environ.get('OPENAI_API_KEY')
 if not key:raise RuntimeError('OPENAI_API_KEY not set')
 prompt={ 'task':'Rewrite the supplied source records as two DRAFT Markdown reports: learner and instructor. Do not invent facts, numerical scores, severity, priority or individual responsibility. Preserve exact source IDs, categories, assessments and quoted evidence. Distinguish evidence from proposed reflection/action. Use Hattie-Timperley goal/current/next and Nicol-Macfarlane-Dick reflection/dialogue. Keep all seven categories visible. Output JSON with learner_markdown and instructor_markdown.', 'records':records[:30], 'targets':TARGETS }
 body={'model':model,'input':[{'role':'system','content':'You are an educational feedback drafting assistant. Clinical correctness and individual attribution require human review. Return valid JSON only.'},{'role':'user','content':json.dumps(prompt,ensure_ascii=False)}]}
 req=urllib.request.Request('https://api.openai.com/v1/responses',data=json.dumps(body).encode(),headers={'Authorization':'Bearer '+key,'Content-Type':'application/json'},method='POST')
 with urllib.request.urlopen(req,timeout=timeout) as f:resp=json.load(f)
 text=''.join(part.get('text','') for o in resp.get('output',[]) for part in o.get('content',[]) if part.get('type')=='output_text')
 text=re.sub(r'^```(?:json)?\s*|\s*```$','',text.strip())
 obj=json.loads(text)
 if not all(k in obj for k in ('learner_markdown','instructor_markdown')):raise ValueError('Missing expected keys')
 return obj

def main():
 p=argparse.ArgumentParser(description=__doc__)
 p.add_argument('workbooks',nargs='+',type=Path)
 p.add_argument('--output',type=Path,default=Path('feedback_output'))
 p.add_argument('--examples-per-category',type=int,default=2)
 p.add_argument('--use-llm',action='store_true',help='Opt-in to API drafting')
 p.add_argument('--allow-external-data',action='store_true',help='Explicitly confirm external transmission is authorised')
 p.add_argument('--model',default=os.environ.get('OPENAI_MODEL','gpt-6'),help='Model ID available to your API account; check current availability')
 p.add_argument('--timeout',type=int,default=90)
 args=p.parse_args()
 if args.use_llm and not args.allow_external_data:p.error('LLM requires --allow-external-data and institutional permission for data transfer')
 args.output.mkdir(parents=True,exist_ok=True)
 records=[]
 for path in args.workbooks:
  if not path.exists():p.error(f'Not found: {path}')
  records+=extract(path)
 (args.output/'source_records.json').write_text(json.dumps(records,indent=2,ensure_ascii=False),encoding='utf-8')
 with (args.output/'audit_template.csv').open('w',newline='',encoding='utf-8') as f:
  w=csv.writer(f);w.writerow(['source_id','evidence_traceability','category_and_polarity','attribution','context_preserved','unsupported_claims','expert_review','notes'])
  for r in records:w.writerow([r['id'],'pending','pending','pending','pending','pending','pending',''])
 build_reports(records,args.output,args.examples_per_category)
 if args.use_llm:
  obj=api_draft(records,args.model,args.timeout)
  for name,key in [('learner_llm_draft.md','learner_markdown'),('instructor_llm_draft.md','instructor_markdown')]:
   (args.output/name).write_text(obj[key],encoding='utf-8')
  (args.output/'llm_run_metadata.json').write_text(json.dumps({'model_requested':args.model,'record_count_sent':min(30,len(records)),'note':'No expert validation performed'},indent=2))
 print(json.dumps({'workbooks':len(args.workbooks),'annotation_records':len(records),'output':str(args.output),'llm':args.use_llm},indent=2))
if __name__=='__main__':main()

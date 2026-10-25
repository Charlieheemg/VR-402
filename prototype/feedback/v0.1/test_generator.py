"""Synthetic regression tests characterize v0.1, including its known defects."""
import subprocess
import gc
import sys
import tempfile
import unittest
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import PatternFill
import fyp_feedback_generator as generator

class FrozenGeneratorTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.cleanup)
        self.root = Path(self.temp.name)
        self.source = self.root / "synthetic.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Summary"
        headers = generator.CATEGORIES.copy()
        headers[5] = "Summarisation and Re-evaluation"
        sheet.append(["Team role"] + headers)
        for row, color in enumerate(["0092D050", "00FF0000", "FF5FA052"], start=2):
            sheet.cell(row, 1, f"Invented training dialogue {row}")
            for column in range(2, 9):
                cell = sheet.cell(row, column, f"Synthetic observation {row}-{column}")
                cell.fill = PatternFill("solid", fgColor=color)
        workbook.save(self.source)
        workbook.close()

    def cleanup(self):
        # Frozen v0.1 leaves workbook handles to garbage collection.
        gc.collect()
        self.temp.cleanup()

    def test_canonical_ampersand_header_is_not_recognized(self):
        self.assertIsNone(generator.normalize("Summarisation & Re-evaluation"))
        self.assertEqual(generator.CATEGORIES[5],
                         generator.normalize("Summarisation and Re-evaluation"))

    def test_extraction_keeps_cell_text_category_and_unknown_attribution(self):
        records = generator.extract(self.source)
        self.assertEqual(21, len(records))
        self.assertEqual(set(generator.CATEGORIES), {r["category"] for r in records})
        self.assertTrue(all(r["attribution"] == "team_or_unverified" for r in records))
        self.assertEqual("Synthetic observation 2-2", records[0]["assessment_text"])
        self.assertEqual("B2", records[0]["cell"])

    def test_known_and_unmapped_colors_remain_distinct(self):
        records = generator.extract(self.source)
        for row, expected in [(2, "positive_color"), (3, "negative_color"), (4, "unverified")]:
            self.assertEqual({expected}, {r["polarity"] for r in records if r["row"] == row})

    def test_default_display_cap_and_repeatability(self):
        records = generator.extract(self.source)
        generator.build_reports(records, self.root, 2)
        learner = (self.root / "learner_draft.md").read_bytes()
        instructor = (self.root / "instructor_draft.md").read_bytes()
        self.assertEqual(14, learner.count(b"**Observed annotation:**"))
        self.assertNotIn(b"Synthetic observation 4-2", learner)
        generator.build_reports(records, self.root, 2)
        self.assertEqual(learner, (self.root / "learner_draft.md").read_bytes())
        self.assertEqual(instructor, (self.root / "instructor_draft.md").read_bytes())

    def test_api_mode_rejects_missing_external_data_flag(self):
        command = [sys.executable, str(Path(generator.__file__)), str(self.source),
                   "--output", str(self.root / "blocked"), "--use-llm"]
        result = subprocess.run(command, capture_output=True, text=True)
        self.assertNotEqual(0, result.returncode)
        self.assertIn("LLM requires --allow-external-data", result.stderr)
        self.assertFalse((self.root / "blocked").exists())

if __name__ == "__main__":
    unittest.main()
